# -*- coding: utf-8 -*-
"""Genere les livrables Excel, Word (liens cliquables) et Markdown depuis final.json.
Sortie: Articles_pertinents.xlsx / .docx / Synthese_articles_pertinents.md"""
import json, collections, datetime

rows = json.load(open("final.json", encoding="utf-8"))
cnt = collections.Counter(r["decision"] for r in rows)
DATE = datetime.date.today().isoformat()
TIERS = [
    ("INCLURE_SOUSNATIONAL", "Categorie A - Donnees sous-nationales d'un seul pays (analogue direct)"),
    ("INCLURE_NATIONAL", "Categorie B - Donnees nationales d'un seul pays"),
    ("BORDERLINE", "Categorie C - A reverifier (cas limites)"),
]
TIERS = [(d, l) for d, l in TIERS if any(r["decision"] == d for r in rows)]

# ===================== EXCEL =====================
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws0 = wb.active; ws0.title = "Synthese"; ws0.sheet_view.showGridLines = False
synth = [
    ("Tri de litterature Scopus - articles pertinents", 16, True),
    ("Genere le " + DATE, 10, False), ("", 10, False),
    ("Articles retenus : %d" % len(rows), 12, True),
    ("  Categorie A (sous-national, un pays) : %d" % cnt.get("INCLURE_SOUSNATIONAL", 0), 11, False),
    ("  Categorie B (national, un pays) : %d" % cnt.get("INCLURE_NATIONAL", 0), 11, False),
    ("  Categorie C (a reverifier) : %d" % cnt.get("BORDERLINE", 0), 11, False),
    ("", 10, False),
    ("Avertissement : tri base sur les abstracts. Verifier le texte integral avant citation.", 10, False),
    ("Colonne 'Lien d'acces' de l'onglet Articles : cliquable (DOI direct ou Google Scholar).", 10, False),
]
for i, (txt, sz, bold) in enumerate(synth, 1):
    ws0.cell(row=i, column=1, value=txt).font = Font(size=sz, bold=bold)
ws0.column_dimensions["A"].width = 110

ws = wb.create_sheet("Articles")
headers = ["Rang", "n", "Categorie", "Score", "Pays", "Niveau geo", "Unite geo",
           "Mesure d'inegalite", "Mesure environnementale", "Pourquoi pertinent",
           "Titre", "Auteurs", "Source", "Volume", "Annee", "Type", "Doublon possible",
           "DOI", "Lien d'acces", "Abstract"]
ws.append(headers)
hfill = PatternFill("solid", fgColor="1F4E78")
thin = Side(style="thin", color="D9D9D9"); border = Border(thin, thin, thin, thin)
for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col); c.font = Font(bold=True, color="FFFFFF", size=10)
    c.fill = hfill; c.alignment = Alignment(vertical="center", wrap_text=True); c.border = border
tier_fill = {"INCLURE_SOUSNATIONAL": PatternFill("solid", fgColor="E2EFDA"),
             "INCLURE_NATIONAL": PatternFill("solid", fgColor="FFF2CC"),
             "BORDERLINE": PatternFill("solid", fgColor="FCE4D6")}
tier_short = {"INCLURE_SOUSNATIONAL": "A. Sous-national", "INCLURE_NATIONAL": "B. National", "BORDERLINE": "C. A reverifier"}
for i, r in enumerate(rows, 1):
    ws.append([i, r["n"], tier_short.get(r["decision"], "?"), r.get("score"), r["country"],
               r["geo_level"], r["geo_unit"], r["inequality_measure"], r["env_measure"],
               r["reason"], r["title"], r["authors"], r["source"], r["volume"], r["year"],
               r["doctype"], r.get("doublon_possible", ""), r.get("doi", ""),
               r["access_link"], r["abstract"]])
    er = i + 1
    lc = ws.cell(row=er, column=19); lc.value = "Ouvrir (%s)" % r.get("link_type", "lien")
    lc.hyperlink = r["access_link"]; lc.font = Font(color="0563C1", underline="single", size=10)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=er, column=col); cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if col == 3: cell.fill = tier_fill.get(r["decision"], PatternFill())
        if col in (1, 2, 4, 15): cell.alignment = Alignment(vertical="top", horizontal="center")
widths = [6, 6, 16, 8, 14, 13, 22, 28, 26, 50, 55, 22, 26, 12, 7, 14, 16, 26, 18, 90]
for col, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(col)].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = "A1:%s1" % get_column_letter(len(headers))
ws.row_dimensions[1].height = 30
wb.save("Articles_pertinents.xlsx")
print("Excel OK : Articles_pertinents.xlsx")

# ===================== WORD =====================
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_hyperlink(p, url, text, color="0563C1", bold=True, size=12):
    r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    col = OxmlElement("w:color"); col.set(qn("w:val"), color); rPr.append(col)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    if bold: rPr.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(size * 2)); rPr.append(sz)
    run.append(rPr); t = OxmlElement("w:t"); t.text = text; run.append(t)
    h.append(run); p._p.append(h); return h

doc = Document(); doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)
doc.add_heading("Articles pertinents - tri Scopus", level=0)
doc.add_paragraph(
    "Genere le %s. %d articles retenus : %d sous-nationaux et %d nationaux. "
    "Le titre de chaque article est cliquable (DOI direct, ou Google Scholar si DOI introuvable). "
    "Avertissement : tri base sur les abstracts, verifier le texte integral avant citation."
    % (DATE, len(rows), cnt.get("INCLURE_SOUSNATIONAL", 0), cnt.get("INCLURE_NATIONAL", 0)))
for dec, label in TIERS:
    sub = [r for r in rows if r["decision"] == dec]
    if not sub: continue
    doc.add_heading("%s  (%d articles)" % (label, len(sub)), level=1)
    for idx, r in enumerate(sub, 1):
        pt = doc.add_paragraph(); pt.add_run("%d. " % idx).bold = True
        add_hyperlink(pt, r["access_link"], r["title"], size=12)
        m = doc.add_paragraph(); m.paragraph_format.space_after = Pt(2)
        run = m.add_run("%s (%s). %s%s" % (r["authors"] or "Auteurs n.d.", r["year"],
                        r["source"], (", " + r["volume"]) if r["volume"] else ""))
        run.font.size = Pt(10); run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        dp = doc.add_paragraph(); dp.paragraph_format.space_after = Pt(2)
        if r.get("doi"):
            dr = dp.add_run("DOI : "); dr.font.size = Pt(9.5); dr.font.color.rgb = RGBColor(0x59,0x59,0x59)
            add_hyperlink(dp, r["access_link"], r["doi"], bold=False, size=9.5)
        else:
            dr = dp.add_run("Lien : "); dr.font.size = Pt(9.5); dr.font.color.rgb = RGBColor(0x59,0x59,0x59)
            add_hyperlink(dp, r["access_link"], "Google Scholar (DOI introuvable)", bold=False, size=9.5)
        info = doc.add_paragraph(); info.paragraph_format.space_after = Pt(2)
        tg = info.add_run("Pays : %s  |  Niveau : %s (%s)  |  Inegalite : %s  |  Environnement : %s" % (
            r["country"] or "n.d.", r["geo_level"] or "n.d.", r["geo_unit"] or "n.d.",
            r["inequality_measure"] or "n.d.", r["env_measure"] or "n.d."))
        tg.font.size = Pt(9.5); tg.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
        wy = doc.add_paragraph(); wy.paragraph_format.space_after = Pt(2)
        wr = wy.add_run("Pertinence : " + (r["reason"] or "")); wr.font.size = Pt(10); wr.italic = True
        if r.get("doublon_possible"):
            dpz = doc.add_paragraph(); z = dpz.add_run("Doublon possible : " + r["doublon_possible"])
            z.font.size = Pt(9); z.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        ab = doc.add_paragraph(); a = ab.add_run(r["abstract"] or "(abstract non disponible)")
        a.font.size = Pt(9.5); ab.paragraph_format.space_after = Pt(10)
doc.save("Articles_pertinents.docx")
print("Word OK : Articles_pertinents.docx")

# ===================== MARKDOWN =====================
L = ["# Articles pertinents - tri Scopus", "",
     "Genere le %s. **%d articles retenus** : %d sous-nationaux et %d nationaux. "
     "Tri base sur les abstracts ; verifier le texte integral avant citation."
     % (DATE, len(rows), cnt.get("INCLURE_SOUSNATIONAL", 0), cnt.get("INCLURE_NATIONAL", 0)), ""]
for dec, label in TIERS:
    sub = [r for r in rows if r["decision"] == dec]
    if not sub: continue
    L.append("## %s (%d)" % (label, len(sub))); L.append("")
    for idx, r in enumerate(sub, 1):
        L.append("### %d. [%s](%s)" % (idx, r["title"].replace("[", "(").replace("]", ")"), r["access_link"]))
        L.append("")
        L.append("- **Auteurs / source** : %s (%s). %s%s" % (
            r["authors"] or "n.d.", r["year"], r["source"], (", " + r["volume"]) if r["volume"] else ""))
        if r.get("doi"):
            L.append("- **DOI** : [%s](%s)" % (r["doi"], r["access_link"]))
        else:
            L.append("- **Lien** : [Google Scholar](%s) (DOI introuvable)" % r["access_link"])
        L.append("- **Pays / niveau** : %s, %s (%s)" % (r["country"] or "n.d.", r["geo_level"] or "n.d.", r["geo_unit"] or "n.d."))
        L.append("- **Inegalite / environnement** : %s ; %s" % (r["inequality_measure"] or "n.d.", r["env_measure"] or "n.d."))
        L.append("- **Pertinence** : %s" % (r["reason"] or ""))
        if r.get("doublon_possible"): L.append("- **Doublon possible** : %s" % r["doublon_possible"])
        L.append(""); L.append("> %s" % (r["abstract"] or "(abstract non disponible)")); L.append("")
open("Synthese_articles_pertinents.md", "w", encoding="utf-8").write("\n".join(L))
print("Markdown OK : Synthese_articles_pertinents.md | total :", len(rows))
